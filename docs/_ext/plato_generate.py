"""Generate the guide's reference material from the normative files at build time.

Writes, under docs/_generated/ (git-ignored):
  sheets.md          column reference for every sheet, from schemas/tables/csv-metadata.json
  vocabularies.md    every concept scheme and RelationType, from ontology.ttl
  downloads/         the template workbook (.xlsx), the template CSVs and each worked
                     example as a zip, all built from schemas/tables/
"""
import csv, io, json, os, re, zipfile
from pathlib import Path

DOCS = Path(__file__).resolve().parent.parent
REPO = DOCS.parent
TABLES = REPO / "schemas" / "tables"
OUT = DOCS / "_generated"
P = "https://w3id.org/plato#"
METADATA_URL = "https://w3id.org/plato/schemas/tables/csv-metadata.json"
NUMERIC = {"decimal", "nonNegativeInteger", "integer"}


def _cell(text):
    return str(text).replace("|", "\\|").replace("\n", " ")


def allowed_values(column):
    dt = column.get("datatype")
    if isinstance(dt, dict):
        m = re.fullmatch(r"\^\(([A-Za-z|]+)\)\$", dt.get("format", ""))
        if m:
            return m.group(1).split("|")
        if dt.get("base") == "boolean" and re.fullmatch(r"[A-Za-z]+\|[A-Za-z]+", dt.get("format", "")):
            return dt["format"].split("|")   # a CSVW boolean written as, say, yes|no
    return None


def value_hint(column):
    dt = column.get("datatype")
    vals = allowed_values(column)
    if vals:
        return "One of: " + ", ".join(f"`{v}`" for v in vals)
    hint = ""
    if dt == "anyURI":
        hint = "A web address"
    elif dt == "nonNegativeInteger":
        hint = "A whole number"
    elif isinstance(dt, dict):
        base = dt.get("base")
        if base == "decimal":
            lo, hi = dt.get("minInclusive"), dt.get("maxInclusive")
            hint = "A number" + (f" from {lo} to {hi}" if lo is not None and hi is not None else f", at least {lo}" if lo is not None else "")
        elif "-?\\d{4" in dt.get("format", ""):
            hint = "Year of at least four digits, or YYYY-MM-DD"
        elif dt.get("@id", "").endswith("wktLiteral"):
            hint = "Well-Known Text"
    if column.get("separator"):
        hint = (hint + "; " if hint else "") + f"several allowed, separated by `{column['separator']}`"
    return hint


def data_columns(table):
    return [c for c in table["tableSchema"]["columns"] if not c.get("virtual")]


def write_sheets_md(meta):
    lines = []
    for t in meta["tables"]:
        name = t["url"][:-4]
        lines += [f"## {name}", "", t["dc:description"], ""]
        lines += ["| Column | Needed? | What to put | Example | Values |", "|---|---|---|---|---|"]
        for c in data_columns(t):
            ex = c.get("skos:example")
            lines.append("| `{}` | {} | {} | {} | {} |".format(
                c["name"], "**required**" if c.get("required") else "optional",
                _cell(c["dc:description"]), f"`{_cell(ex)}`" if ex else "", _cell(value_hint(c))))
        lines.append("")
    (OUT / "sheets.md").write_text("\n".join(lines), encoding="utf-8")


def write_vocabularies_md():
    from rdflib import Graph, Namespace, RDF, RDFS
    SKOS = Namespace("http://www.w3.org/2004/02/skos/core#")
    PL = Namespace(P)
    g = Graph()
    g.parse(REPO / "ontology.ttl", format="turtle")
    lines = []
    for scheme in sorted(g.subjects(RDF.type, SKOS.ConceptScheme), key=str):
        label = g.value(scheme, RDFS.label)
        lines += [f"## {str(label).capitalize()}", "", str(g.value(scheme, RDFS.comment)), "",
                  "| Value | Meaning | Full identifier |", "|---|---|---|"]
        for c in sorted(g.subjects(SKOS.inScheme, scheme), key=str):
            lines.append("| `{}` | {} | `{}` |".format(str(c)[len(P):], _cell(g.value(c, SKOS.definition)), c))
        lines.append("")
    lines += ["## Certainty levels", "",
              "The words that the `certainty_level` column, and `certaintyLevel` in JSON, can name, for certainty given in words rather than as a number.",
              "", "| Value | Meaning | Full identifier |", "|---|---|---|"]
    for c in sorted(g.subjects(RDF.type, PL.CertaintyLevel), key=str):
        if str(c).startswith(P):
            lines.append("| `{}` | {} | `{}` |".format(str(c)[len(P):], _cell(g.value(c, RDFS.comment)), c))
    lines.append("")
    lines += ["## Relation types", "",
              "The relations that the `relation_type` column of the relations sheet, and `relationType` in JSON, can name.",
              "", "| Value | Reads as | Inverse | Meaning |", "|---|---|---|---|"]
    for r in sorted(g.subjects(RDF.type, PL.RelationType), key=str):
        lines.append("| `{}` | {} | {} | {} |".format(str(r)[len(P):], _cell(g.value(r, PL.relation_label)),
                     _cell(g.value(r, PL.inverse_label)), _cell(" ".join(str(g.value(r, RDFS.comment)).split()))))
    lines += ["", "## Routes, itineraries, networks and segments", "",
              "Types that platforms act on: they draw a route from its ordered members, and may leave segments out of lists of places. "
              "Put the full identifier in the `type_uri` column of the types sheet, or `identifier` of a type in JSON, with the value as the `type_label`.",
              "", "| Value | Meaning | Full identifier |", "|---|---|---|"]
    for t in sorted(g.subjects(RDF.type, PL.Type), key=str):
        if str(t).startswith(P):
            lines.append("| `{}` | {} | `{}` |".format(g.value(t, PL.type_label), _cell(g.value(t, RDFS.comment)), t))
    (OUT / "vocabularies.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_workbook(meta, path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    rows = [
        ["temPlato: the PLATO spreadsheet template"],
        [""],
        ["Fill in one sheet per kind of information. Keep all ten sheets, even the ones you leave empty."],
        ["Hover over a column heading to see what to put in it and an example."],
        ["Every row in names, locations, types, relations, connections and properties needs a place_id (from the places sheet), a source_id (from the sources sheet) and a date."],
        ["Write the date as your source gives it, or 'undated'. Put years in 'from' and 'to' with at least four digits: 0921, not 921."],
        ["Columns with a fixed list of values offer a drop-down; the values are case-sensitive."],
        [""],
        ["Guide: https://pelagios.org/place-attestation-ontology/guide/"],
        ["Table definitions (CSVW): " + METADATA_URL],
    ]
    for r in rows:
        readme.append(r)
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 120
    head_fill = PatternFill("solid", fgColor="DDE7F0")
    # Excel refuses an inline drop-down list longer than 255 characters, and the workbook then
    # opens as damaged; the CiTO functions are far longer. Every list therefore lives in a column
    # of a hidden sheet, which a drop-down can reference at any length.
    lists = wb.create_sheet("values")
    lists.sheet_state = "hidden"
    list_col = 0
    for t in meta["tables"]:
        ws = wb.create_sheet(t["url"][:-4])
        cols = data_columns(t)
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=i, value=c["name"])
            cell.font = Font(bold=True)
            cell.fill = head_fill
            note = c["dc:description"] + (f"\n\nExample: {c['skos:example']}" if c.get("skos:example") else "")
            note += "\n\n" + ("Required." if c.get("required") else "Optional.")
            cell.comment = Comment(note, "PLATO", width=320, height=180)
            letter = cell.column_letter
            ws.column_dimensions[letter].width = max(12, min(40, len(c["name"]) + 6))
            dt = c.get("datatype")
            base = dt.get("base") if isinstance(dt, dict) else dt
            if base not in NUMERIC:
                # Text format stops spreadsheet programs turning 0921 into 921 or 1-2 into a date.
                for r in range(2, 1002):
                    ws.cell(row=r, column=i).number_format = "@"
            vals = allowed_values(c)
            if vals:
                list_col += 1
                for r, v in enumerate(vals, start=1):
                    lists.cell(row=r, column=list_col, value=v)
                ref = lists.cell(row=1, column=list_col).column_letter
                dv = DataValidation(type="list", formula1=f"values!${ref}$1:${ref}${len(vals)}", allow_blank=True)
                listed = ", ".join(vals)
                # Excel's error message holds at most 225 characters.
                dv.error = ("Choose one of: " + listed) if len(listed) <= 200 else "Choose a value from the drop-down list; the values are case-sensitive."
                dv.errorTitle = "Not an allowed value"
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}1001")
        ws.freeze_panes = "A2"
    wb.move_sheet(lists, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("values"))   # last, out of the way
    wb.save(path)


def write_zip(path, csv_dir):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(TABLES / "csv-metadata.json", "csv-metadata.json")
        for f in sorted(Path(csv_dir).glob("*.csv")):
            z.write(f, f.name)


def write_example_tables():
    """One markdown table per non-empty sheet of each worked example, showing only
    the columns that have a value somewhere, so the walkthroughs quote the real files."""
    for ex in sorted((TABLES / "examples").iterdir()):
        d = OUT / "examples" / ex.name
        d.mkdir(parents=True, exist_ok=True)
        for f in sorted(ex.glob("*.csv")):
            rows = list(csv.DictReader(f.open(encoding="utf-8")))
            if not rows:
                (d / (f.stem + ".md")).write_text("*(empty apart from its heading row)*\n", encoding="utf-8")
                continue
            cols = [c for c in rows[0].keys() if any(r[c] for r in rows)]
            lines = ["| " + " | ".join(f"`{c}`" for c in cols) + " |", "|" + "---|" * len(cols)]
            lines += ["| " + " | ".join(_cell(r[c]) for c in cols) + " |" for r in rows]
            (d / (f.stem + ".md")).write_text("\n".join(lines) + "\n", encoding="utf-8")


def generate(app=None):
    OUT.mkdir(exist_ok=True)
    (OUT / "downloads").mkdir(exist_ok=True)
    meta = json.loads((TABLES / "csv-metadata.json").read_text(encoding="utf-8"))
    write_sheets_md(meta)
    write_vocabularies_md()
    write_example_tables()
    write_workbook(meta, OUT / "downloads" / "plato-tables-template.xlsx")
    write_zip(OUT / "downloads" / "plato-tables-template.zip", TABLES)
    for ex in sorted((TABLES / "examples").iterdir()):
        write_zip(OUT / "downloads" / f"plato-tables-example-{ex.name}.zip", ex)


def setup(app):
    app.connect("builder-inited", generate)
    return {"parallel_read_safe": True, "parallel_write_safe": True}
